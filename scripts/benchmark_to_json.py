#!/usr/bin/env python3
"""
全行业财务指标基准表 Excel → JSON 转换脚本

用法:
    pip install openpyxl
    python scripts/benchmark_to_json.py

输入: Excel 文件（默认从 rag-data/staging/ 或命令行参数读取）
输出: rag-data/processed/national/rates/industry_benchmark.json

JSON 结构:
    {
      "meta": { "source": "...", "generated_at": "...", "total_industries": 97 },
      "industries": [
        {
          "category": "农、林、牧、渔业",
          "sub_industry": "农业",
          "vat_burden": {"low": 0.00, "high": 0.02},
          "cit_burden": {"low": 0.00, "high": 0.05},
          "gross_margin": {"low": 0.20, "high": 0.35},
          "net_margin": {"low": 0.05, "high": 0.12},
          "ar_turnover": {"low": 3, "high": 8},
          "inventory_turnover": {"low": 2, "high": 6},
          "debt_ratio": {"low": 0.40, "high": 0.60},
          "current_ratio": {"low": 1.0, "high": 2.0},
          "quick_ratio": {"low": 0.5, "high": 1.2},
          "expense_ratio": {"low": 0.08, "high": 0.15}
        },
        ...
      ]
    }
"""

import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("错误: 需要安装 openpyxl")
    print("  pip install openpyxl")
    sys.exit(1)

# ── 路径配置 ──────────────────────────────────────────
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
OUTPUT_DIR = PROJECT_ROOT / "rag-data" / "processed" / "national" / "rates"
STAGING_DIR = PROJECT_ROOT / "rag-data" / "staging"

# 列映射 (Excel 列序号 → 字段名)
# A=1: 国家行业分类（门类）
# B=2: 核心细分行业
# C=3: 平均增值税税负范围
# D=4: 平均企业所得税税负范围
# E=5: 行业平均毛利率
# F=6: 行业平均净利率
# G=7: 应收账款周转率（次/年）
# H=8: 存货周转率（次/年）
# I=9: 行业资产负债率
# J=10: 行业流动比率
# K=11: 行业速动比率
# L=12: 行业费用率范围
COLUMN_MAP = {
    3:  "vat_burden",
    4:  "cit_burden",
    5:  "gross_margin",
    6:  "net_margin",
    7:  "ar_turnover",
    8:  "inventory_turnover",
    9:  "debt_ratio",
    10: "current_ratio",
    11: "quick_ratio",
    12: "expense_ratio",
}

# 中文标签映射（用于前端展示）
LABEL_MAP = {
    "vat_burden":         "增值税税负率",
    "cit_burden":         "企业所得税税负率",
    "gross_margin":       "毛利率",
    "net_margin":         "净利率",
    "ar_turnover":        "应收账款周转率",
    "inventory_turnover": "存货周转率",
    "debt_ratio":         "资产负债率",
    "current_ratio":      "流动比率",
    "quick_ratio":        "速动比率",
    "expense_ratio":      "费用率",
}


def parse_range(text: str):
    """解析 '0.00-0.02' → {"low": 0.00, "high": 0.02}
    解析 '--' → None
    解析 '3-8' → {"low": 3, "high": 8}（整数）
    解析 '0.5-2' → {"low": 0.5, "high": 2}（浮点数）
    """
    if not text or text.strip() in ("--", "", "N/A"):
        return None

    text = text.strip()

    # 处理 "0.00-0.02" 格式
    if "-" in text:
        parts = text.split("-")
        if len(parts) == 2:
            low_str = parts[0].strip()
            high_str = parts[1].strip()
            try:
                # 判断是整数还是浮点数
                def parse_num(s):
                    s = s.strip()
                    if "." in s:
                        return float(s)
                    else:
                        return int(s)

                low = parse_num(low_str)
                high = parse_num(high_str)
                return {"low": low, "high": high}
            except ValueError:
                return None

    return None


def find_excel_file() -> Path:
    """查找 Excel 文件"""
    # 1. 命令行参数
    if len(sys.argv) > 1:
        p = Path(sys.argv[1])
        if p.exists():
            return p

    # 2. staging 目录
    if STAGING_DIR.exists():
        for f in STAGING_DIR.glob("*.xlsx"):
            if "财务指标" in f.name or "benchmark" in f.name.lower():
                return f
        # 如果只有一个 xlsx，直接用
        xlsx_files = list(STAGING_DIR.glob("*.xlsx"))
        if len(xlsx_files) == 1:
            return xlsx_files[0]

    print(f"错误: 未找到 Excel 文件")
    print(f"  请将文件放入 {STAGING_DIR}/ 或通过参数指定:")
    print(f"  python scripts/benchmark_to_json.py <path/to/file.xlsx>")
    sys.exit(1)


def convert(excel_path: Path) -> dict:
    """读取 Excel 并转换为结构化 JSON"""
    wb = openpyxl.load_workbook(str(excel_path), data_only=True)
    ws = wb[wb.sheetnames[0]]

    industries = []
    current_category = None

    for row in ws.iter_rows(min_row=3, values_only=False):
        # 读取门类（合并单元格可能只有第一行有值）
        cat_val = row[0].value
        if cat_val and cat_val.strip():
            current_category = cat_val.strip()

        sub_val = row[1].value
        if not sub_val or not sub_val.strip():
            continue

        entry = {
            "category": current_category,
            "sub_industry": sub_val.strip(),
        }

        for col_idx, field_name in COLUMN_MAP.items():
            cell_val = row[col_idx - 1].value
            if isinstance(cell_val, str):
                entry[field_name] = parse_range(cell_val)
            elif cell_val is not None:
                entry[field_name] = {"low": float(cell_val), "high": float(cell_val)}
            else:
                entry[field_name] = None

        industries.append(entry)

    wb.close()

    result = {
        "meta": {
            "source": "全行业财务指标基准表.xlsx",
            "source_path": str(excel_path),
            "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "total_industries": len(industries),
            "field_labels": LABEL_MAP,
            "notes": "数值为行业平均范围，low=下限，high=上限。null 表示该指标不适用于此行业（如金融业的毛利率）。",
        },
        "industries": industries,
    }

    return result


def main():
    excel_path = find_excel_file()
    print(f"读取 Excel: {excel_path}")

    result = convert(excel_path)

    # 确保输出目录存在
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_DIR / "industry_benchmark.json"

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"输出 JSON: {output_path}")
    print(f"共 {result['meta']['total_industries']} 个细分行业")

    # 统计行业分布
    from collections import Counter
    cat_counts = Counter(ind["category"] for ind in result["industries"])
    print(f"\n行业门类分布:")
    for cat, count in cat_counts.most_common():
        print(f"  {cat}: {count} 个细分")

    # 统计 null 值
    null_stats = {}
    for field in COLUMN_MAP.values():
        null_count = sum(1 for ind in result["industries"] if ind[field] is None)
        if null_count > 0:
            null_stats[field] = null_count
    if null_stats:
        print(f"\n含 null 值的指标:")
        for field, count in null_stats.items():
            print(f"  {LABEL_MAP.get(field, field)}: {count} 个行业不适用")


if __name__ == "__main__":
    main()

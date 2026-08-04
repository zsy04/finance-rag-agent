#!/usr/bin/env python3
"""
个税申报表填表工具 — Agent @tool 函数
=======================================
由 Agent 调用，将对话中收集的用户信息填入申报表模板。

用法（Agent 侧）：
  from scripts.fill_tax_form import fill_form

  result = fill_form(
      form_type="A表",
      user_data={
          "纳税人姓名": "张三",
          "身份证件类型": "居民身份证",
          "身份证件号码": "4101xxxxxxxxxxxxxx",
          ...
      },
      output_path="/tmp/申报表_张三.xlsx"
  )
"""

import json
import shutil
from pathlib import Path
from datetime import datetime

try:
    import openpyxl
except ImportError:
    openpyxl = None


# 模板和映射文件路径（相对于项目根目录）
PROJECT_ROOT = Path(__file__).resolve().parent.parent
TEMPLATES_DIR = PROJECT_ROOT / "rag-data" / "processed" / "national" / "templates"
FIELD_MAP_FILE = TEMPLATES_DIR / "form_field_map.json"


def _load_field_map():
    """加载字段映射配置"""
    with open(FIELD_MAP_FILE, "r", encoding="utf-8") as f:
        return json.load(f)


def _parse_cell(cell_ref):
    """将 'C9' 转成 (row, col)，col 从1开始"""
    col_letter = ""
    row_num = ""
    for ch in cell_ref:
        if ch.isalpha():
            col_letter += ch
        else:
            row_num += ch
    # 列字母转数字：A=1, B=2, ..., Z=26, AA=27
    col = 0
    for ch in col_letter.upper():
        col = col * 26 + (ord(ch) - ord('A') + 1)
    return int(row_num), col


def fill_form(form_type: str, user_data: dict, output_path: str = None) -> str:
    """
    将用户数据填入申报表模板，返回填写完成的文件路径。

    Args:
        form_type: 表单类型 "A表" 或 "B表"
        user_data: 字段名 → 值的字典。字段名与 form_field_map.json 中的 key 对应。
        output_path: 输出路径。默认输出到项目 outputs/ 目录。

    Returns:
        JSON 字符串: {"success": true/false, "file": "路径", "message": "..."}

    Examples:
        >>> result = fill_form("A表", {"纳税人姓名": "张三", "身份证件号码": "4101..."})
        >>> print(result)
        {"success": true, "file": "/path/to/申报表_张三.xlsx", ...}
    """
    if openpyxl is None:
        return json.dumps({
            "success": False,
            "message": "openpyxl 未安装，请运行 pip install openpyxl"
        }, ensure_ascii=False)

    # 加载映射
    try:
        mapping = _load_field_map()
    except Exception as e:
        return json.dumps({
            "success": False,
            "message": f"加载字段映射失败: {e}"
        }, ensure_ascii=False)

    if form_type not in mapping["forms"]:
        return json.dumps({
            "success": False,
            "message": f"未知表单类型: {form_type}。可选: {list(mapping['forms'].keys())}"
        }, ensure_ascii=False)

    form_config = mapping["forms"][form_type]
    template_file = TEMPLATES_DIR / form_config["file"]

    if not template_file.exists():
        return json.dumps({
            "success": False,
            "message": f"模板文件不存在: {template_file}"
        }, ensure_ascii=False)

    # 创建输出目录
    name = user_data.get("纳税人姓名", "未命名")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    if not output_path:
        out_dir = PROJECT_ROOT / "outputs"
        out_dir.mkdir(exist_ok=True)
        output_path = out_dir / f"申报表_{name}_{timestamp}.xlsx"
    else:
        output_path = Path(output_path)

    # 复制模板
    shutil.copy(template_file, output_path)

    # 填写表单
    try:
        from openpyxl.comments import Comment

        wb = openpyxl.load_workbook(output_path)
        ws = wb[form_config.get("sheet", wb.sheetnames[0])]

        filled_count = 0
        skipped = []
        errors = []
        notes_items = []  # 收集所有备注字段

        for field_name, value in user_data.items():
            if field_name not in form_config["fields"]:
                skipped.append(field_name)
                continue

            field_config = form_config["fields"][field_name]
            cell_ref = field_config["cell"]

            # 备注字段：不填格子，收集起来统一写到备注区域
            if cell_ref == "_备注_":
                if value and value != "0":
                    note_text = field_config.get("note", field_name)
                    notes_items.append(f"{note_text}: {value}")
                    filled_count += 1
                continue

            try:
                row, col = _parse_cell(cell_ref)
                cell = ws.cell(row=row, column=col)
                cell.value = value
                filled_count += 1
            except Exception as e:
                errors.append(f"{field_name} → {cell_ref}: {e}")

        # 将经营数据作为备注写入表格末尾（第 36 行附近）
        if notes_items:
            try:
                note_row = 36
                ws.cell(row=note_row, column=1, value="【经营所得申报数据】")
                ws.cell(row=note_row, column=1).font = openpyxl.styles.Font(bold=True)
                for i, item in enumerate(notes_items):
                    ws.cell(row=note_row + 1 + i, column=1, value=item)
                filled_count += 1  # 备注区标题也算一次填写
            except Exception as e:
                errors.append(f"备注写入: {e}")

        wb.save(output_path)

        message_parts = [f"已填写 {filled_count} 个字段"]
        if skipped:
            message_parts.append(f"跳过未知字段: {', '.join(skipped)}")
        if errors:
            message_parts.append(f"错误: {'; '.join(errors)}")

        return json.dumps({
            "success": True,
            "file": str(output_path.absolute()),
            "form_type": form_type,
            "filled_fields": filled_count,
            "skipped_fields": skipped,
            "message": " | ".join(message_parts)
        }, ensure_ascii=False)

    except Exception as e:
        return json.dumps({
            "success": False,
            "message": f"填表失败: {e}"
        }, ensure_ascii=False)


def get_required_fields(form_type: str) -> str:
    """
    获取指定表单的必填字段列表，供 Agent 判断还需收集哪些信息。

    Returns:
        JSON 字符串: {"required": [...], "optional": [...]}
    """
    try:
        mapping = _load_field_map()
    except Exception as e:
        return json.dumps({"error": str(e)}, ensure_ascii=False)

    if form_type not in mapping["forms"]:
        return json.dumps({"error": f"未知表单: {form_type}"}, ensure_ascii=False)

    fields = mapping["forms"][form_type]["fields"]
    required = []
    optional = []

    for name, config in fields.items():
        item = {
            "field": name,
            "ask": config.get("ask", ""),
            "type": config.get("type", "text"),
            "options": config.get("options"),
            "default": config.get("default")
        }
        if config.get("required"):
            required.append(item)
        else:
            optional.append(item)

    return json.dumps({
        "form_type": form_type,
        "required": required,
        "optional": optional
    }, ensure_ascii=False)


# ── Agent 工具调用示例 ──
if __name__ == "__main__":
    # 示例：填写 A表
    sample_data = {
        "扣缴义务人名称": "郑州科技有限公司",
        "扣缴义务人识别号": "91410100MA12345678",
        "纳税人姓名": "张三",
        "身份证件类型": "居民身份证",
        "身份证件号码": "410105199001011234",
        "出生日期": "1990-01-01",
        "国籍/地区": "中国",
        "任职受雇从业类型": "雇员",
        "职务": "工程师",
        "学历": "大学本科",
        "任职受雇从业日期": "2020-03-01",
        "手机号码": "13800138000",
        "联系地址": "郑州市金水区XX路XX号",
        "是否残疾/孤老/烈属": "否"
    }

    result = fill_form("A表", sample_data)
    print(result)
